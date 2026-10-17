import io
import zipfile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from tests.media.conftest import image_bytes, jobs, mclient, storage  # noqa: F401  (fixtures + helper)


def add_link(paragraph, url: str, text: str) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def build_docx(*, image: bytes | None = None, bad_link: bool = False) -> bytes:
    d = Document()
    d.add_heading("Flood Report", level=1)
    d.add_heading("Details", level=2)
    p = d.add_paragraph("Plain then ")
    p.add_run("bold").bold = True
    p.add_run(" and ")
    p.add_run("italic").italic = True
    p.add_run(" and ")
    p.add_run("under").underline = True
    p.add_run(" and ")
    p.add_run("struck").font.strike = True
    p.add_run(" see ")
    add_link(p, "https://example.com/a", "the site")
    if bad_link:
        p.add_run(" and ")
        add_link(p, "javascript:alert(1)", "evil")
    d.add_paragraph("first item", style="List Bullet")
    d.add_paragraph("second item", style="List Bullet")
    d.add_paragraph("step one", style="List Number")
    d.add_paragraph("Wise words", style="Quote")
    t = d.add_table(rows=2, cols=2)
    for i, row in enumerate(t.rows):
        for j, cell in enumerate(row.cells):
            cell.text = f"r{i}c{j}"
    if image is not None:
        d.add_picture(io.BytesIO(image), width=Pt(100))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def docx_files(data: bytes, name="report.docx") -> dict:
    return {"file": (name, data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}


def zip_with(members: dict[str, bytes], compression=zipfile.ZIP_DEFLATED) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression) as z:
        for name, data in members.items():
            z.writestr(name, data)
    return buf.getvalue()
